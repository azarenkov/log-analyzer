import sys, pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
OUT=sys.argv[1]
summ=pd.read_csv("out/results_summary.csv"); log=open("out/run_log.txt").read(); code=open("pilot_pipeline.py").read()
d=Document(); st=d.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(12)
for s in d.sections: s.left_margin=s.right_margin=Cm(2.5)
H=lambda t,l=1:d.add_heading(t,l)
def P(t,b=False,i=False):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=b; r.italic=i; p.paragraph_format.space_after=Pt(6); return p
def T(rows,hdr,w=None,fs=9):
    t=d.add_table(rows=1,cols=len(hdr)); t.style="Table Grid"
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i]; c.text=""; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(fs)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=""; r=cs[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(fs)
    d.add_paragraph(); return t
def MONO(t):
    p=d.add_paragraph(); r=p.add_run(t); r.font.name="Courier New"; r.font.size=Pt(7.5)

tp=d.add_paragraph(); tp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tp.add_run("Assignment 3. Methodology and Core Practical Implementation"); r.bold=True; r.font.size=Pt(16)
P("Students: Rauan Kuntuganov, Alexey Azarenkov"); P("Group: SE-2401")
P("Core track: Machine Learning / Deep Learning / AI Model")

H("1 Final Research Title")
P("Cross-System Generalization of Log Anomaly Detectors on Loghub Benchmarks (unchanged from Assignment 2).")

H("2 Primary Research Question, Aim and Objectives")
P("Primary research question: How does anomaly-detection F1 differ among selected classical, recurrent and Transformer model representatives across within-system training, zero-shot cross-system transfer and adaptation using 5% of the target training partition on HDFS, BGL and Thunderbird?")
P("Aim: To evaluate how anomaly-detection F1 differs among selected classical, recurrent and Transformer model representatives across within-system training, zero-shot transfer and 5% target adaptation on HDFS, BGL and Thunderbird.")
for i,o in enumerate([
 "Identify one reproducible model representative from each of the classical, recurrent and Transformer categories, and assess compatibility with the three datasets and declared training-data conditions.",
 "Define comparable dataset partitions, evaluation units, label handling and target-data access rules, including a fixed 5% adaptation budget and a held-out target test partition.",
 "Evaluate each selected representative under within-system training, zero-shot transfer and 5% target adaptation for all six directed source-to-target dataset pairs.",
 "Compare F1 differences between settings and representatives, supported by precision, recall, false-positive rate and training cost, and assess variation across repeated runs."],1):
    P(f"O{i}. {o}")
P("Hypotheses (from Assignment 2): H1 – mean F1 under zero-shot transfer is lower than under within-system training; H2 – mean F1 after 5% adaptation is higher than under zero-shot transfer. The research question, aim and objectives are unchanged.")

H("3 Research Approach and Design")
P("The study is quantitative and computational. Its design is a controlled comparative benchmark experiment with two crossed factors: model representative (classical, recurrent, Transformer) and training/transfer setting (within-system, zero-shot, 5% adaptation), repeated over the six directed pairs formed by HDFS, BGL and Thunderbird. Everything except the manipulated factor is held constant: the target test partition, the log representation, the label rules and the metric code.")
P("The unit of analysis is one experimental run: a model–setting–direction combination evaluated on a fixed target test partition and summarised by F1. Inside each run the prediction unit is a log sequence (an HDFS block session or a fixed window of BGL/Thunderbird lines), because the Loghub labels are attached at that level.")
P("The design fits the question because the question asks how a numeric outcome differs across named conditions; only a controlled experiment isolates the effect of the setting from differences in data splits or evaluation code, which Le and Zhang (2022) showed can change reported scores. It produces quantitative evidence: F1, precision, recall, false-positive rate (FPR) and training time per run, with variation across seeds.")
P("Rejected alternative: a secondary comparison of scores already published, for example the multi-dataset results of DeepLog, LogBERT, SemPCA or the cross-system LogGT model (Du et al., 2017; Guo et al., 2021; Yang et al., 2024; Wang et al., 2024). It is cheaper, but those studies use different splits, grouping, supervision and target-data access, so the differences would mix model effects with protocol effects; this is exactly the comparability gap identified in Assignment 1.")

H("4 Dataset and Research Context")
P("Research material. The population is labelled system logs from distributed and supercomputer systems; the sample is three public Loghub datasets (Zhu et al., 2023): HDFS (Hadoop file system, labels per block, Xu et al., 2009), BGL (BlueGene/L supercomputer, labels per line, Oliner & Stearley, 2007) and Thunderbird (supercomputer, labels per line). No human participants are involved.")
P("Inclusion criteria: publicly released Loghub datasets with ground-truth anomaly labels and pre-parsed templates; log lines with a parsable message. Exclusion criteria: lines without a block identifier (HDFS) or without content; Loghub datasets without labels (e.g., most Loghub 2k sets other than BGL/Thunderbird); target validation or test data in any transfer training step.")
P("Size and rationale. The full study uses chronological subsets: all of HDFS (≈575,000 sessions) and the first 5 million lines of BGL and Thunderbird, a bounded size that fits a student GPU budget within one semester while keeping both normal and anomalous periods. Each dataset is split chronologically 70/10/20 (train/validation/test); the 5% adaptation budget is drawn only from the target training partition. The practical pilot in Section 9 used smaller public samples (Table 2) because the full archives are hosted on Zenodo and had to be downloaded separately.")
P("Selection strategy. Purposive dataset selection: these three are the most frequently used labelled benchmarks in the reviewed studies (He et al., 2016; Le & Zhang, 2022), which makes results comparable to prior work. Access is realistic: all files are openly available from Loghub/Zenodo and GitHub without a licence request.")
P("Selection bias. Benchmarks are cleaner and better documented than production logs, and anomalies are clustered in time, so a random split leaks future patterns into training. We reduce this by chronological splitting, by reporting the anomaly ratio of every partition, and by checking (as the pilot did) that every test partition actually contains anomalies before using it.")

H("5 Variables, Measures and Metrics")
P("Independent variables: (1) model representative – a supervised classical model on semantic template features (logistic regression; He et al., 2016), kept because SemPCA showed that lightweight semantic classical detectors can compete with neural ones (Yang et al., 2024); a LogRobust-style attention Bi-LSTM over semantic template vectors (recurrent; Zhang et al., 2019); and a LogFormer-style Transformer encoder over the same vectors (Guo et al., 2024). All three are supervised, so differences are not caused by different supervision: Assignment 1 showed that self-supervised LogBERT and semi-supervised PLELog receive different label information (Guo et al., 2021; Yang et al., 2021), which is why they were not chosen; (2) setting – within-system, zero-shot, 5% adaptation; (3) source→target direction (six levels).")
P("Dependent variable: F1 for the anomaly class on the target test partition (primary). Supporting outcomes: precision, recall, FPR = FP/(FP+TN), and wall-clock training time.")
P("Controlled variables: same test partition across settings, same window size (20 lines) and session grouping, same representation family (words of parsed templates, so that vocabularies can overlap across systems; DeepLog-style event IDs are treated as unrelated symbols and cannot carry meaning to a new system, Du et al., 2017; Meng et al., 2019), fixed decision threshold 0.5 unless tuned inside the budget, fixed random seeds, same library versions.")
P("Instruments are software: Loghub ground-truth labels, the pre-parsed Drain templates shipped with Loghub, scikit-learn metric functions (Pedregosa et al., 2011) and a run logger that writes every configuration and result to CSV. F1, precision and recall answer O3–O4; partition statistics and anomaly counts answer O2; the feasibility checks of each model with each dataset answer O1.")

H("6 Experimental Procedure")
P("1) Preparation: obtain the Loghub files, record checksums, fix the environment (Python, scikit-learn, PyTorch versions in a requirements file) and seeds 0–4. 2) Preprocessing: group HDFS lines by block ID and join with anomaly_label.csv; cut BGL and Thunderbird into non-overlapping chronological windows of 20 lines, a window being anomalous if any line is labelled. 3) Representation: normalise each template to lowercase words (placeholders and numbers removed). 4) Partitioning: chronological 70/10/20 split per dataset; log the size and anomaly ratio of each part. 5) Within-system runs: train each model on target train, select hyper-parameters on target validation, test on target test. 6) Zero-shot runs: train on source train/validation only; the vectoriser and threshold are fitted on source data; predict the target test partition unchanged. 7) Adaptation runs: draw 5% of target training units with seeds 0–4 and update the source model with them (fine-tuning for neural models, refitting on source+adaptation data for the classical model); nothing else from the target is seen. 8) Repeat every run for five seeds, store raw predictions and metrics.")
P("Fairness: all settings for one target share the identical test file; all models receive the same inputs; the adaptation sample is identical across models for a given seed. Raw data are kept read-only; derived partitions are written to separate files named by dataset, split and seed.")

H("7 Data Analysis Plan")
P("O1 – Feasibility: a compatibility checklist per model and dataset (runs end-to-end, memory, training time, whether a single-class partition blocks training).")
P("O2 – Comparability: descriptive tables of units, anomalies and anomaly ratio per partition and per adaptation sample; vocabulary overlap between datasets.")
P("O3 – Evaluation: for each of 3 models × 3 settings × 6 directions, mean and standard deviation of F1, precision, recall and FPR over five seeds; confusion matrices.")
P("O4 – Comparison and hypotheses: for H1 and H2 the paired units are the 18 model–direction combinations (seed-averaged). We report the mean paired difference with a bootstrap 95% confidence interval and a one-sided Wilcoxon signed-rank test (α = 0.05), recommended for comparing classifiers over several datasets (Demšar, 2006). Model differences inside a setting are described with mean ranks rather than claimed as family-level conclusions, consistent with Assignment 1's finding that no universal deep-versus-classical ranking is supported (Yang et al., 2024; Yu et al., 2024). Error analysis inspects false positives/negatives to separate vocabulary mismatch from label-definition differences (block vs. line labels). Training cost is compared descriptively.")
P("Leakage prevention: chronological splits, vectoriser fitted only on data the setting is allowed to see, no target test or validation labels in zero-shot or adaptation.")

H("8 Research Quality, Ethics, Privacy and Reproducibility")
P("Validity: F1 on the anomaly class is used instead of accuracy because anomalies are rare (≈3.9% of HDFS sessions in the pilot). Internal validity is protected by identical test partitions and the leakage rules above; construct validity is limited by the different label definitions across datasets, which we report rather than hide. Reliability: five seeds, reported standard deviation.")
P("Reproducibility: public data with recorded sources and SHA-256 checksums; one script (pilot_pipeline.py) regenerates all pilot tables, and fig.py and build_report.py regenerate the figure and this report; pinned versions in requirements.txt; fixed seeds; raw per-run results (results_runs.csv), aggregated results (results_summary.csv) and the unedited run log are saved with the code. The complete package is submitted as a Git repository (Appendix A) so that the instructor can rerun the pilot with: pip install -r requirements.txt; python pilot_pipeline.py data out. Results with Thunderbird as target are reported as N/A for F1 rather than 0, because F1 is undefined without positive examples.")
P("Ethics and privacy: no human participants, so informed consent is not required. Logs may still contain host names, user names (e.g., 'root', node IDs) and IP addresses; we use only the published anonymised benchmarks, do not attempt re-identification, do not publish raw log excerpts beyond short template examples, and keep the data on our own machines. No log data are uploaded to external AI or cloud services; AI tools were used only for language checking, and every methodological decision was verified by the authors. Data are deleted from personal machines after the course, while scripts and aggregated results are kept.")

H("9 Core Practical Component: Initial Implementation")
P("Supports: O1 (one representative per category runs end-to-end on the same inputs), O2 (chronological 70/10/20 partitions, label units and the 5% budget are implemented) and O3 (first local, zero-shot and 5% adaptation runs). Track: ML model pipeline. All numbers below are copied from the files written by the script (results_runs.csv, results_summary.csv, run_log.txt); nothing was edited by hand.")
P("Data. Real Loghub samples: HDFS_100k.log_structured.csv with anomaly_label.csv (logpai/loglizer repository) and BGL_2k / Thunderbird_2k log_structured files (logpai/loghub repository). SHA-256 prefixes of the files are printed in the run log. The full Loghub archives (Zenodo) will replace these samples in Assignment 4.")
P("Models (same input for all three: each event is the list of words of its parsed template). Classical: logistic regression on TF-IDF of template words (class_weight=balanced). Recurrent: LogRobust-style attention Bi-LSTM; each event vector is the mean of learned word embeddings (dim 32), Bi-LSTM hidden size 32, attention pooling. Transformer: 2-layer encoder, 4 heads, learned positional embeddings over the same event vectors. Neural models: Adam (lr 0.002), batch 64, positive-class weight = normal/anomaly ratio, 8 epochs with early stopping on validation F1 of data the setting is allowed to see (target validation for local runs, source validation for zero-shot); adaptation fine-tunes the source model for 5 epochs on the 5% sample only, with the source vocabulary (unseen words become UNK). Logistic regression is deterministic and run once; neural models use seeds 0, 1, 2. Environment: Python 3.13, PyTorch 2.14.1 (CPU), scikit-learn 1.9.1, pandas 3.0.5, numpy 2.5.3 (requirements.txt).")
s=pd.read_csv("out/dataset_summary.csv")
P("Table 1. Pilot partitions after preprocessing (dataset_summary.csv). HDFS unit = block session; BGL/Thunderbird unit = window of 20 lines.",b=True)
T([[r.dataset,f"{r.train} / {r.train_anom}",f"{r.val} / {r.val_anom}",f"{r.test} / {r.test_anom}",r.templates] for r in s.itertuples()],["Dataset","Train (units / anomalies)","Validation","Test","Templates"])
P("Table 2. Pilot results on the target test partition (results_summary.csv; mean ± SD over seeds for neural models). Thunderbird targets are N/A for F1, precision and recall because the test partition contains no anomalies; only FPR is defined there.",b=True)
names={"logreg":"LogReg","bilstm":"Bi-LSTM","transformer":"Transformer"}
fmt=lambda v,sd=None: "N/A" if pd.isna(v) else (f"{v:.3f}"+(f" ± {sd:.3f}" if sd is not None and pd.notna(sd) else ""))
rows=[[names[r.model],r.setting,(f"{r.source}→{r.target}" if r.source!=r.target else r.source),fmt(r.F1_mean,r.F1_sd),fmt(r.P_mean),fmt(r.R_mean),fmt(r.FPR_mean),f"{r.train_s:.1f}"] for r in summ.itertuples()]
T(rows,["Model","Setting","Direction","F1","Precision","Recall","FPR","Train s"],fs=8)
d.add_picture("out/f1_pilot.png",width=Cm(16))
P("Figure 1. F1 on the target test partition for directions with defined F1 (error bars: SD over seeds). Thunderbird targets are excluded because F1 is N/A.",i=True)
P("What the pilot shows (descriptive only; no hypothesis test is applied to pilot data). Local F1 on HDFS was 0.699 (LogReg), 0.693 ± 0.011 (Bi-LSTM) and 0.595 ± 0.066 (Transformer). Zero-shot BGL→HDFS was low for all models (0.278, 0.109, 0.096), and the 5% HDFS sample (278 sessions, 6–16 anomalies) raised F1 for all three (0.562, 0.466 ± 0.147, 0.461 ± 0.311). This is the direction predicted by H1 and H2, but with six usable comparisons and two tiny BGL partitions it is not evidence for either hypothesis.")
P("Decision explained: EventIds such as 'E5' are labels assigned separately by the parser for each dataset, so 'E5' in HDFS and 'E5' in BGL are unrelated and cannot transfer; this confirms in our own data the concern from Assignment 1 (Meng et al., 2019). All three models therefore read the words of templates, which share 11–68 words between dataset pairs (run log). We also remove numbers and placeholders from templates; NeuralLog warns that such cleaning can delete informative codes (Le & Zhang, 2021), which we record as a known trade-off.")
P("Failures and unresolved issues found:",b=True)
for x in [
 "Thunderbird_2k contains no labelled anomalies. Local training and transfer from Thunderbird were skipped automatically, and every result with Thunderbird as target is reported as N/A for F1 and excluded from the comparison; only FPR is shown there. The full study must choose a Thunderbird slice with anomalies in every partition, checked before training.",
 "HDFS→BGL zero-shot logistic regression labelled all 12 normal BGL test windows as anomalous (FP = 12, TN = 0, FPR = 1.0) and found 6 of 8 anomalous windows (FN = 2, recall = 0.750). Its F1 of 0.462 therefore comes mostly from the high anomaly share in the BGL test part (8 of 20) and must be read together with FPR. Several neural zero-shot runs (e.g., BGL→HDFS) also labelled every target unit as anomalous.",
 "The 5% budget on BGL is only 4 windows, and for all seeds this sample contained no anomalies. Fine-tuning on four normal windows pushed the neural models to predict 'normal' for everything (Transformer HDFS→BGL adapt F1 = 0.000 in all seeds), so 'adaptation' made them worse. With full datasets the budget becomes thousands of units; we will also report the anomaly count of every adaptation sample and decide in Assignment 4 whether stratified sampling within the budget is acceptable.",
 "Neural results vary strongly between seeds on small partitions (e.g., Bi-LSTM local BGL F1 = 0.485 ± 0.241), and the BGL validation set has only 10 windows, so early stopping is unstable. Larger data and more seeds are needed.",
 "Earlier test: restricting HDFS to its first 4,000 sessions made local logistic regression collapse (F1 = 0.077, FPR = 0.873), while all 7,940 sessions gave F1 = 0.699. Small chronological subsets can therefore be unrepresentative; the full study will check partition statistics before training."]:
    d.add_paragraph(x,style="List Bullet")
P("Next steps: download the full Loghub archives, select subsets with anomalies in every partition, keep the same script and configuration, increase epochs and seeds, and run the statistical analysis of Section 7 only on the full data.")

H("10 Methodology Alignment Table")
T([
 ["O1 Select representatives","Runnable implementation per category; compatibility record","Published methods (Yang 2024; Zhang 2019; Guo 2024); Loghub","Pilot runs, run log, environment file","Compatibility checklist; training time","Confirms the comparison in the RQ can be executed"],
 ["O2 Comparable partitions & budgets","Unit counts, anomaly ratios, budget sizes, vocabulary overlap","HDFS, BGL, Thunderbird (Loghub)","Preprocessing script; chronological 70/10/20 split","Descriptive statistics per partition","Ensures differences are due to settings, not data"],
 ["O3 Evaluate 3 settings × 6 directions","Predictions on target test for every run","Target test partitions","Training/evaluation script, 5 seeds","P, R, F1, FPR, confusion matrix; mean ± SD","Produces the F1 values the RQ asks about"],
 ["O4 Compare & test H1/H2","Paired F1 differences; cost","Results CSV from O3","Aggregation script","Wilcoxon signed-rank, bootstrap CI, ranks, error analysis","Answers how F1 differs and tests H1/H2"]],
 ["Objective","Data / evidence needed","Source / dataset","Collection / instrument","Analysis method","Evidence for RQ"],fs=8)

H("11 References")
REFS=[
 "Demšar, J. (2006). Statistical comparisons of classifiers over multiple data sets. Journal of Machine Learning Research, 7, 1–30.",
 'Du, M., Li, F., Zheng, G., & Srikumar, V. (2017). DeepLog: Anomaly detection and diagnosis from system logs through deep learning. Proceedings of the 2017 ACM SIGSAC Conference on Computer and Communications Security, 1285–1298. https://doi.org/10.1145/3133956.3134015',
 'Guo, H., Yuan, S., & Wu, X. (2021). LogBERT: Log anomaly detection via BERT. 2021 International Joint Conference on Neural Networks. https://doi.org/10.1109/IJCNN52387.2021.9534113',
 'Le, V. H., & Zhang, H. (2021). Log-based anomaly detection without log parsing. 2021 36th IEEE/ACM International Conference on Automated Software Engineering, 492–504. https://doi.org/10.1109/ASE51524.2021.9678773',
 'Meng, W., Liu, Y., Zhu, Y., Zhang, S., Pei, D., Liu, Y., Chen, Y., Zhang, R., Tao, S., Sun, P., & Zhou, R. (2019). LogAnomaly: Unsupervised detection of sequential and quantitative anomalies in unstructured logs. Proceedings of the Twenty-Eighth International Joint Conference on Artificial Intelligence, 4739–4745. https://doi.org/10.24963/ijcai.2019/658',
 'Wang, P., Zhang, X., Cao, Z., Xu, W., & Li, W. (2024). LogGT: Cross-system log anomaly detection via heterogeneous graph feature and transfer learning. Expert Systems with Applications, 251, 124082. https://doi.org/10.1016/j.eswa.2024.124082',
 'Yang, L., Chen, J., Gao, S., Gong, Z., Zhang, H., Kang, Y., & Li, H. (2024). Try with simpler — An evaluation of improved principal component analysis in log-based anomaly detection. ACM Transactions on Software Engineering and Methodology. https://doi.org/10.1145/3644386',
 'Yang, L., Chen, J., Wang, Z., Wang, W., Jiang, J., Dong, X., & Zhang, W. (2021). Semi-supervised log-based anomaly detection via probabilistic label estimation. 2021 IEEE/ACM 43rd International Conference on Software Engineering, 1448–1460. https://doi.org/10.1109/ICSE43902.2021.00130',
 'Yu, B., Yao, J., Fu, Q., Zhong, Z., Xie, H., Wu, Y., Ma, Y., & He, P. (2024). Deep learning or classical machine learning? An empirical study on log-based anomaly detection. Proceedings of the IEEE/ACM 46th International Conference on Software Engineering, 1–13. https://doi.org/10.1145/3597503.3623308',
 "Guo, H., Yang, J., Liu, J., Bai, J., Wang, B., Li, Z., Zheng, T., Zhang, B., Peng, J., & Tian, Q. (2024). LogFormer: A pre-train and tuning pipeline for log anomaly detection. Proceedings of the AAAI Conference on Artificial Intelligence, 38(1), 135–143. https://doi.org/10.1609/aaai.v38i1.27764",
 "He, S., Zhu, J., He, P., & Lyu, M. R. (2016). Experience report: System log analysis for anomaly detection. 2016 IEEE 27th International Symposium on Software Reliability Engineering, 207–218. https://doi.org/10.1109/ISSRE.2016.21",
 "Le, V. H., & Zhang, H. (2022). Log-based anomaly detection with deep learning: How far are we? Proceedings of the 44th International Conference on Software Engineering, 1356–1367. https://doi.org/10.1145/3510003.3510155",
 "Oliner, A., & Stearley, J. (2007). What supercomputers say: A study of five system logs. 37th Annual IEEE/IFIP International Conference on Dependable Systems and Networks, 575–584. https://doi.org/10.1109/DSN.2007.103",
 "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
 "Xu, W., Huang, L., Fox, A., Patterson, D., & Jordan, M. I. (2009). Detecting large-scale system problems by mining console logs. Proceedings of the ACM SIGOPS 22nd Symposium on Operating Systems Principles, 117–132. https://doi.org/10.1145/1629575.1629587",
 "Zhang, X., Xu, Y., Lin, Q., Qiao, B., Zhang, H., Dang, Y., Xie, C., Yang, X., Qian, C., Li, Z., Chen, J., He, X., Yao, R., Lou, J. G., Chintalapati, M., Shen, F., & Zhang, D. (2019). Robust log-based anomaly detection on unstable log data. Proceedings of ESEC/FSE 2019, 807–817. https://doi.org/10.1145/3338906.3338931",
 "Zhu, J., He, S., He, P., Liu, J., & Lyu, M. R. (2023). Loghub: A large collection of system log datasets for AI-driven log analytics. 2023 IEEE 34th International Symposium on Software Reliability Engineering, 355–366. https://doi.org/10.1109/ISSRE59848.2023.00071"]
if 1:
    pass
refs=sorted(REFS)
for ref in refs:
    p=d.add_paragraph(ref); p.paragraph_format.left_indent=Cm(1); p.paragraph_format.first_line_indent=Cm(-1)

H("12 Appendix")
H("Appendix A. Data sources and repository contents",2)
for x in ["https://github.com/logpai/loghub – BGL/BGL_2k.log_structured.csv, Thunderbird/Thunderbird_2k.log_structured.csv",
          "https://github.com/logpai/loglizer – data/HDFS/HDFS_100k.log_structured.csv, data/HDFS/anomaly_label.csv",
          "Full datasets for Assignment 4: Loghub on Zenodo (Zhu et al., 2023).",
 "Repository (Git, submitted with this report): pilot_pipeline.py, fig.py, build_report.py, requirements.txt, README.md, out/dataset_summary.csv, out/results_runs.csv, out/results_summary.csv, out/run_log.txt, out/f1_pilot.png. Add the GitHub link here before submission: ______________________"]:
    d.add_paragraph(x,style="List Bullet")
H("Appendix B. Run log (unedited console output of pilot_pipeline.py)",2); MONO(log)
H("Appendix C. Pilot pipeline source code (pilot_pipeline.py)",2); MONO(code)
H("Appendix D. Optional bonus",2)
P("The optional survey/interview bonus was not attempted: the study has no human-participant component, and expert feedback would not add evidence about F1 differences.")
d.save(OUT)
